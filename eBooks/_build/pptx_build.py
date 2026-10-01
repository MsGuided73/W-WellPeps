"""Build an editable PowerPoint from pptx_extract.mjs output: one portrait slide per eBook page, every box, rule,
photo, icon and text block as its own PowerPoint object at the position, size, colour and type the PDF uses.

Usage: uv run --with python-pptx --with fonttools python pptx_build.py pptx/<book>.json out.pptx

Text keeps the browser's line breaks (soft line breaks inside each text box) so it wraps exactly as the PDF does;
fonts are Inter / Inter Medium / Inter SemiBold / Lora SemiBold (desktop files in pptx/fonts/). Photos keep their
crop as an editable PowerPoint crop, not a pre-cut image.
"""
import base64
import io
import json
import math
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU = 9525  # per CSS px (96 dpi)
# PowerPoint places the first baseline of an exact-spaced line a little lower than CSS does; measured against the PDF
BASELINE_NUDGE = 0.0
WIDEN = 1.04  # text boxes get 4% extra width so PowerPoint never wraps a line the browser kept whole
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def e(v):
    return Emu(int(round(v * EMU)))


def parse_color(s):
    """'rgb(r, g, b)' / 'rgba(r, g, b, a)' -> (RGBColor, alpha 0..1) or (None, 0)."""
    if not s:
        return None, 0
    m = re.match(r'rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)', s)
    if not m:
        return None, 0
    r, g, b = (int(float(m.group(i))) for i in (1, 2, 3))
    a = float(m.group(4)) if m.group(4) is not None else 1.0
    return RGBColor(r, g, b), a


def set_alpha(color_elm, alpha):
    if alpha < 0.999:
        a = etree.SubElement(color_elm, qn('a:alpha'))
        a.set('val', str(int(alpha * 100000)))


def solid_fill(shape, css):
    col, a = parse_color(css)
    shape.fill.solid()
    shape.fill.fore_color.rgb = col
    set_alpha(shape.fill._xPr.find(qn('a:solidFill'))[0], a)


def gradient_fill(shape, css, w):
    """CSS linear-gradient -> DrawingML gradFill (angle, stops with alpha)."""
    body = css[css.index('(') + 1: css.rindex(')')]
    parts = re.split(r',\s*(?![^()]*\))', body)
    angle = 180.0
    if parts and re.match(r'^-?[\d.]+deg$', parts[0].strip()):
        angle = float(parts[0].strip()[:-3]); parts = parts[1:]
    stops = []
    for p in parts:
        m = re.match(r'\s*(rgba?\([^)]*\))\s*([\d.]+)(%|px)?', p)
        if not m:
            continue
        pos = float(m.group(2))
        pos = pos / 100 if m.group(3) == '%' else pos / w
        stops.append((min(max(pos, 0), 1), m.group(1)))
    spPr = shape._element.spPr
    for tag in ('a:noFill', 'a:solidFill', 'a:gradFill'):
        for old in spPr.findall(qn(tag)):
            spPr.remove(old)
    grad = etree.Element(qn('a:gradFill'), rotWithShape='1')
    lst = etree.SubElement(grad, qn('a:gsLst'))
    for pos, c in stops:
        col, a = parse_color(c)
        gs = etree.SubElement(lst, qn('a:gs'), pos=str(int(pos * 100000)))
        clr = etree.SubElement(gs, qn('a:srgbClr'), val=str(col))
        set_alpha(clr, a)
    # CSS 180deg runs top->bottom, which is DrawingML 90deg; CSS 90deg (left->right) is DrawingML 0
    etree.SubElement(grad, qn('a:lin'), ang=str(int(((angle - 90) % 360) * 60000)), scaled='0')
    geom = spPr.find(qn('a:prstGeom'))
    geom.addnext(grad)


def no_line(shape):
    shape.line.fill.background()


def hard_shadow(shape, css):
    """Print-media shadow '2px 3px 0 #d9dee7' -> opaque offset outer shadow (editable in Format Shape)."""
    m = re.match(r'(rgba?\([^)]*\))\s+(-?[\d.]+)px\s+(-?[\d.]+)px\s+([\d.]+)px', css)
    if not m:
        return
    col, a = parse_color(m.group(1)); dx, dy, blur = float(m.group(2)), float(m.group(3)), float(m.group(4))
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad=str(int(blur * EMU)), dist=str(int(math.hypot(dx, dy) * EMU)),
                          dir=str(int(math.degrees(math.atan2(dy, dx)) * 60000)), algn='tl', rotWithShape='0')
    clr = etree.SubElement(sh, qn('a:srgbClr'), val=str(col))
    set_alpha(clr, a)


def round_adj(shape, r, w, h):
    shape.adjustments[0] = min(0.5, r / max(1.0, min(w, h)))


def plain(shape):
    """Drop the default theme style (python-pptx autoshapes reference the theme's effect style, which draws a soft
    shadow under every shape in PowerPoint); fill and line are always set explicitly here."""
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)
    return shape


def add_box(slide, it):
    x, y, w, h = it['x'], it['y'], it['w'], it['h']
    if w <= 0 or h <= 0:
        return
    sides = it['sides']
    live = [s for s in sides if s['w'] > 0 and s['s'] != 'none' and parse_color(s['c'])[1] > 0]
    uniform = len(live) == 4 and len({(round(s['w'], 2), s['c']) for s in live}) == 1
    radius = max(it['radius'])
    if it['fill'] or it['grad'] or it['shadow'] or uniform:
        shp = plain(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, e(x), e(y), e(w), e(h)))
        if radius:
            round_adj(shp, radius, w, h)
        if it['grad']:
            gradient_fill(shp, it['grad'], w)
        elif it['fill']:
            solid_fill(shp, it['fill'])
        else:
            shp.fill.background()
        if uniform:
            col, a = parse_color(live[0]['c'])
            shp.line.color.rgb = col; shp.line.width = Pt(live[0]['w'] * 0.75)
        else:
            no_line(shp)
        if it['shadow']:
            hard_shadow(shp, it['shadow'])
        shp.name = (it.get('cls') or 'box').split(' ')[0] or 'box'
        shp.text_frame.text = ''
    if not uniform:  # single-side borders (the callout's blue bar, hairline rules) as thin rectangles
        for side, s in zip(('top', 'right', 'bottom', 'left'), sides):
            if s['w'] <= 0 or s['s'] == 'none' or parse_color(s['c'])[1] == 0:
                continue
            bw = s['w']
            rect = {'top': (x, y, w, bw), 'bottom': (x, y + h - bw, w, bw), 'left': (x, y, bw, h), 'right': (x + w - bw, y, bw, h)}[side]
            bar = plain(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(e(v) for v in rect)))
            solid_fill(bar, s['c']); no_line(bar); bar.name = f'{side} border'


def add_image(slide, it, images):
    src = images[it['img']]
    raw = base64.b64decode(src.split(',', 1)[1])
    x, y, w, h = it['x'], it['y'], it['w'], it['h']
    nat_r = it['natW'] / it['natH'] if it['natH'] else w / h
    box_r = w / h
    crop = [0, 0, 0, 0]  # l, t, r, b
    if it['fit'] == 'cover' and abs(nat_r - box_r) > 1e-3:
        def frac(v, excess_px):  # CSS object-position component -> 0..1 share of the overflow cut from the left/top
            if v.endswith('%'):
                return float(v[:-1]) / 100
            if v.endswith('px'):
                return min(1.0, max(0.0, float(v[:-2]) / excess_px)) if excess_px > 0.5 else 0.5
            return {'left': 0, 'top': 0, 'center': 0.5, 'right': 1, 'bottom': 1}.get(v, 0.5)
        pos = (it['pos'].split() + ['50%'])[:2]
        exw = w * (nat_r / box_r) - w if box_r < nat_r else 0
        exh = h * (box_r / nat_r) - h if box_r > nat_r else 0
        px_, py_ = frac(pos[0], exw), frac(pos[1], exh)
        if box_r > nat_r:  # photo taller than the slot: crop top/bottom
            ex = 1 - nat_r / box_r; crop[1], crop[3] = ex * py_, ex * (1 - py_)
        else:
            ex = 1 - box_r / nat_r; crop[0], crop[2] = ex * px_, ex * (1 - px_)
    elif it['fit'] == 'contain' and abs(nat_r - box_r) > 1e-3:
        if box_r > nat_r:
            nw = h * nat_r; x += (w - nw) / 2; w = nw
        else:
            nh = w / nat_r; y += (h - nh) / 2; h = nh
    pic = slide.shapes.add_picture(io.BytesIO(raw), e(x), e(y), e(w), e(h))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    radii = it.get('radii') or [it.get('radius', 0)] * 4
    r = max(radii) or it.get('radius', 0)
    if r:
        geom = pic._element.spPr.find(qn('a:prstGeom'))
        top_only = radii[2] == 0 and radii[3] == 0 and radii[0] > 0
        geom.set('prst', 'round2SameRect' if top_only else 'roundRect')
        av = geom.find(qn('a:avLst'))
        if av is None:
            av = etree.SubElement(geom, qn('a:avLst'))
        etree.SubElement(av, qn('a:gd'), name='adj1' if top_only else 'adj', fmla=f'val {int(min(50000, r / min(w, h) * 100000))}')
        if top_only:
            etree.SubElement(av, qn('a:gd'), name='adj2', fmla='val 0')
    pic.name = 'icon' if it.get('icon') else 'photo'


# Invisible formatting characters the browser uses (the renderer glues "NAD+—" with a word joiner, U+2060) but the
# embedded desktop fonts do not contain. PowerPoint on a machine without Inter/Lora installed can only draw from the
# embedded subsets, and a character missing from them made it substitute fonts across the whole deck (Healthy Aging
# v7, 2026-09-30). Line breaks are already explicit in the export, so these characters do nothing here: drop them.
INVISIBLE = dict.fromkeys(map(ord, '⁠​‌‍﻿­'))


def face(style):
    fam = style['family'].split(',')[0].strip().strip('"\'')
    wt = int(style['weight'])
    if fam.lower().startswith('lora'):
        return ('Lora SemiBold', False) if 500 <= wt < 700 else ('Lora', wt >= 700)
    if wt >= 700:
        return 'Inter', True
    if wt >= 600:
        return 'Inter SemiBold', False
    if wt >= 500:
        return 'Inter Medium', False
    return 'Inter', False


def add_text(slide, it):
    runs = it['runs']
    if not runs:
        return
    lines = 1 + sum(1 for r in runs if r.get('br'))
    lh = it['lh']
    w = it['w'] * WIDEN + 2
    x = it['x']
    if it['align'] == 'center':
        x -= (w - it['w']) / 2
    elif it['align'] in ('right', 'end'):
        x -= (w - it['w'])
    h = max(it['h'], lines * lh)
    tb = slide.shapes.add_textbox(e(x), e(it['y'] + BASELINE_NUDGE * lh), e(w), e(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY}.get(it['align'], PP_ALIGN.LEFT)
    p.line_spacing = Pt(lh * 0.75)
    p.space_before = p.space_after = Pt(0)
    first_text = ''
    for r in runs:
        if r.get('br'):
            p.add_line_break(); continue
        st = r['style']
        run = p.add_run(); run.text = r['text'].translate(INVISIBLE)
        first_text = first_text or r['text']
        name, bold = face(st)
        f = run.font
        f.name = name; f.size = Pt(round(st['size'] * 0.75 * 2) / 2); f.bold = bold; f.italic = st['italic']
        f.underline = st.get('underline', False)
        col, _ = parse_color(st['color']); f.color.rgb = col
        rPr = run._r.get_or_add_rPr()
        if st['ls']:
            rPr.set('spc', str(int(round(st['ls'] * 0.75 * 100))))
        for tag in ('a:latin', 'a:ea', 'a:cs'):  # make the face stick for every script
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set('typeface', name)
    tb.name = (first_text.translate(INVISIBLE)[:40] or 'text')


def check_glyphs(d):
    """Every character must exist in the desktop font it will be drawn with, or PowerPoint substitutes fonts on
    machines that rely on the embedded copies. Fails the export with the offending characters."""
    import os
    from fontTools.ttLib import TTFont
    here = os.path.dirname(os.path.abspath(__file__))
    files = {'Inter': 'Inter-Regular', 'Inter Medium': 'Inter-Medium', 'Inter SemiBold': 'Inter-SemiBold', 'Lora SemiBold': 'Lora-SemiBold', 'Lora': 'Lora-SemiBold'}
    cmaps = {}
    missing = {}
    for pg in d['pages']:
        for it in pg['items']:
            if it['t'] != 'text':
                continue
            for r in it['runs']:
                if r.get('br'):
                    continue
                name, bold = face(r['style'])
                fn = 'Inter-Bold' if (name == 'Inter' and bold) else files[name]
                if fn not in cmaps:
                    cmaps[fn] = TTFont(os.path.join(here, 'pptx', 'fonts', fn + '.ttf')).getBestCmap()
                for ch in r['text'].translate(INVISIBLE):
                    if ord(ch) not in cmaps[fn] and not ch.isspace():
                        missing.setdefault(f'U+{ord(ch):04X} {ch!r}', set()).add(fn)
    if missing:
        raise SystemExit('characters missing from the embedded fonts: ' + '; '.join(f'{k} in {sorted(v)}' for k, v in missing.items()))


def build(src, out):
    d = json.load(open(src, encoding='utf8'))
    check_glyphs(d)
    prs = Presentation()
    W, H = d['pages'][0]['w'], d['pages'][0]['h']
    prs.slide_width, prs.slide_height = e(W), e(H)
    blank = prs.slide_layouts[6]
    for pg in d['pages']:
        s = prs.slides.add_slide(blank)
        for it in pg['items']:
            {'box': add_box, 'text': add_text}.get(it['t'], lambda sl, x: add_image(sl, x, d['images']))(s, it)
    # Embed only the characters in use, like the known-good decks. Whatever PowerPoint embeds is replaced afterwards by
    # pptx_fonts.py: fresh embeds from 2026-09-30 on (subset or full) drew scrambled letters in PowerPoint for the web.
    prs.part._element.set('saveSubsetFonts', '1')
    prs.save(out)
    print(f'{len(d["pages"])} slides -> {out}')


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
